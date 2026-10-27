from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pathlib import Path
out=Path(__file__).resolve().parent/'BHI_Poster_Draft.pptx'
prs=Presentation();prs.slide_width=Inches(16);prs.slide_height=Inches(9)
s=prs.slides.add_slide(prs.slide_layouts[6]); navy=RGBColor(18,43,67); teal=RGBColor(0,127,140); dark=RGBColor(30,43,54); gray=RGBColor(88,104,118); white=RGBColor(255,255,255)
def shape(x,y,w,h,fill):
 a=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));a.fill.solid();a.fill.fore_color.rgb=fill;a.line.fill.background();return a
def txt(t,x,y,w,h,size=15,bold=False,color=dark):
 a=s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h));tf=a.text_frame;tf.word_wrap=True;tf.margin_left=Inches(.06);tf.margin_top=Inches(.03)
 for j,line in enumerate(t.split('\n')):
  p=tf.paragraphs[0] if j==0 else tf.add_paragraph();p.text=line;p.font.size=Pt(size);p.font.bold=bold if j==0 else False;p.font.color.rgb=color;p.space_after=Pt(7)
 return a
shape(0,0,16,1.45,navy)
txt('BIOFUEL HEALTH INDEX & COMMINGLING FRAMEWORK',.6,.18,14.8,.58,27,True,white)
txt('Engineering proof of concept • P Saashkintan • SIT / PSA Marine context',.65,.92,14.6,.36,13,False,RGBColor(209,231,235))
for x in (.6,5.6,10.6):shape(x,1.7,4.75,5.82,RGBColor(239,245,247))
txt('01  PROBLEM & AIM',.86,1.94,4.2,.38,18,True,teal)
txt('Biofuel quality can change during storage and handling. Repeated topping-up complicates tracking of residual fuel history.\n\nAim: develop a transparent Python tool that combines fuel-quality indicators with recursive effective-age bookkeeping.',.86,2.48,4.15,3.7,15)
txt('02  ENGINEERING METHOD',5.86,1.94,4.2,.38,18,True,teal)
txt('• Literature-informed hazard identification\n• FMEA: six illustrative failure modes\n• MCDA: editable parameter ratings and weights\n• BHI: water + TAN + oxidation stability\n• Separate reference-limit breach check\n• Recursive volume-weighted fuel age',5.86,2.48,4.15,4.3,15)
txt('03  TESTING & LIMITATIONS',10.86,1.94,4.15,.38,18,True,teal)
txt('16 / 16 Python tests passed\n\nIllustrative age example: 6.0 days after replacing 800 L of 1,000 L.\n\nNo operational fuel samples. FMEA and MCDA ratings and BHI penalties are provisional. No fuel safety certification is implied.',10.86,2.48,4.13,4.45,15)
shape(.6,7.68,14.75,.62,navy)
txt('NEXT: supervisor review → verified fuel-grade limits → laboratory comparison',.84,7.82,14.3,.35,17,True,white)
txt('Sources: McCormick & Moriarty (NREL, 2023); ISO 8217:2024; MPA Singapore biofuel bunkering guidance; IEC 60812:2018',.65,8.46,14.65,.32,10,False,gray)
prs.save(out);print(out)
